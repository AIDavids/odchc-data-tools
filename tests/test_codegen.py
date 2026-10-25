import datetime as dt
import io
import sqlite3
import threading

import pytest
from openpyxl import load_workbook

from lookup_tool import codegen
from lookup_tool.codegen import ALPHABET, check_char, connect, export_excel, generate_code, verify_code

NOW = dt.datetime(2026, 10, 20, 10, 0)


def make(con, orin, name="ADA JOHN", who="Tester", now=NOW):
    return generate_code(con, orin=orin, client_name=name, source_file="f", requested_by=who, now=now)


def test_generated_codes_verify_and_sequence(tmp_path):
    con = connect(str(tmp_path / "c.db"))
    a, b = make(con, "111-1111-111"), make(con, "222-2222-222")
    assert a["code"].startswith("AUTH-2610-0001-") and b["code"].startswith("AUTH-2610-0002-")
    assert verify_code(a["code"])[0] and verify_code(b["code"])[0]


def test_same_enrollee_gets_the_same_code_not_a_duplicate(tmp_path):
    con = connect(str(tmp_path / "c.db"))
    first, again = make(con, "111-1111-111"), make(con, "111-1111-111")
    assert again["already_issued"] and again["code"] == first["code"]
    assert con.execute("SELECT COUNT(*) FROM authorization_codes").fetchone()[0] == 1


def test_sequence_restarts_each_month(tmp_path):
    con = connect(str(tmp_path / "c.db"))
    make(con, "1")
    nov = make(con, "2", now=dt.datetime(2026, 11, 2))
    assert nov["code"].startswith("AUTH-2611-0001-")


def test_every_single_digit_error_and_adjacent_swap_is_caught():
    payload = "2610" + "0042"
    good = f"AUTH-2610-0042-{check_char(payload)}"
    assert verify_code(good)[0]
    for i in range(len(payload)):
        for d in "0123456789":
            if d != payload[i]:
                bad = payload[:i] + d + payload[i + 1:]
                if 1 <= int(bad[2:4]) <= 12:  # keep it a valid month so only the check char can object
                    assert not verify_code(f"AUTH-{bad[:4]}-{bad[4:]}-{check_char(payload)}")[0]
    for i in range(len(payload) - 1):
        if payload[i] != payload[i + 1]:
            sw = list(payload); sw[i], sw[i + 1] = sw[i + 1], sw[i]; sw = "".join(sw)
            if 1 <= int(sw[2:4]) <= 12:
                assert not verify_code(f"AUTH-{sw[:4]}-{sw[4:]}-{check_char(payload)}")[0]


@pytest.mark.parametrize("bad", ["", "AUTH-2610-42-X", "AUTH-2613-0001-0", "HELLO", None])
def test_malformed_codes_rejected(bad):
    assert verify_code(bad) == (False, "bad_format")


def test_requested_by_is_required(tmp_path):
    con = connect(str(tmp_path / "c.db"))
    with pytest.raises(ValueError):
        make(con, "1", who="  ")


def test_concurrent_requests_never_share_a_code(tmp_path):
    path = str(tmp_path / "c.db")
    connect(path).close()
    out, errors = [], []

    def worker(i):
        try:
            con = connect(path)
            con.execute("PRAGMA busy_timeout = 10000")
            out.append(make(con, f"orin-{i}")["code"])
            con.close()
        except Exception as e:  # pragma: no cover
            errors.append(e)

    threads = [threading.Thread(target=worker, args=(i,)) for i in range(12)]
    [t.start() for t in threads]; [t.join() for t in threads]
    assert not errors and len(set(out)) == 12


def test_excel_export(tmp_path):
    con = connect(str(tmp_path / "c.db"))
    make(con, "1", name="ADA JOHN"); make(con, "2", name="BOLA KEMI")
    buf = io.BytesIO()
    assert export_excel(con, buf) == 2
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    assert [c.value for c in ws[1]][:2] == ["code", "client_name"] and ws.max_row == 3
