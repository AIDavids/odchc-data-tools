"""Generate FAKE monthly facility workbooks that copy the *mess* of the real ones
(shifted header rows, title rows, header spelling drift, numeric phone numbers that lose
their leading zero, mixed date formats, cover sheets). No real person or record is used.
"""
import argparse
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook

FIRST = ["ADE", "BOLA", "CHIDI", "DAYO", "EMEKA", "FOLA", "GBENGA", "HALIMA", "IFEOMA", "JIDE",
         "KEMI", "LANRE", "MARYAM", "NKEM", "OLU", "PEJU", "QUDUS", "RUKKY", "SEUN", "TOPE",
         "UCHE", "VICTOR", "WALE", "YEMI", "ZAINAB", "TUNDE", "SADE", "KUNLE", "BISI", "FEMI"]
SURN = ["AKINTOLA", "BELLO", "CHUKWU", "DADA", "EZE", "FASHOLA", "GARBA", "HASSAN", "IBRAHIM",
        "JOLAOSO", "KOLAWOLE", "LAWAL", "MUSA", "NWOSU", "OGUNDELE", "PETERS", "QUADRI",
        "RAJI", "SALAMI", "TAIWO", "UMAR", "VINCENT", "WILLIAMS", "YUSUF", "ZUBAIRU"]
FACILITIES = ["GH Riverside", "SSH Hillcrest", "MCH Lakeview", "UTH Central", "GH Eastgate",
              "SSH Northfield", "MCH Southbank", "GH Westmoor", "UTH Annex"]
MONTH_NAMES = ["January", "February", "March", "April", "May", "June", "July", "August",
               "September", "October", "November", "December"]
CLIENT_TYPES = ["ANC", "PNC", "CHILD"]
HEADER_VARIANTS = {
    "S/N": ["S/N", "S/N", "SN"],
    "Name of client": ["Name of client", "NAME OF CLIENT", "Name of Client "],
    "Date of booking (dd-mm-yyyy)": ["Date of booking (dd-mm-yyyy)", "Date of Booking"],
    "ORIN ": ["ORIN ", "ORIN", "Orin"],
    "Place of Issue": ["Place of Issue", "Place of issue"],
    "Hospital number": ["Hospital number", "Hospital No"],
    "Age": ["Age", "AGE"],
    "Address": ["Address", "ADDRESS"],
    "Phone number": ["Phone number", "Phone Number", "PHONE NO"],
    "Type of client": ["Type of client", "Type of Client"],
    "Parity": ["Parity", "PARITY"],
}
DEMO_PEOPLE = ["SAMPLE PATIENT ALPHA", "SAMPLE PATIENT BETA"]  # planted so the demo has a known hit


def fake_name(rng):
    parts = [rng.choice(SURN), rng.choice(FIRST), rng.choice(FIRST)]
    if rng.random() < 0.5:
        parts = parts[:2]
    if rng.random() < 0.25:
        parts = parts[::-1]
    name = " ".join(parts)
    r = rng.random()
    return name.title() if r < 0.2 else (name + " " if r < 0.3 else name)


def fake_date(rng, year, month, style):
    d = dt.date(year, month, rng.randint(1, 28))
    if style == "native":
        return dt.datetime(d.year, d.month, d.day)
    return d.strftime("%d-%m-%Y") if style == "dash" else d.strftime("%d/%m/%Y")


def make_row(rng, sn, facility, year, month, phone_as_number, extra_name=None):
    style = rng.choices(["native", "dash", "slash"], [0.5, 0.3, 0.2])[0]
    phone = f"{rng.choice(['70', '80', '81', '90', '91'])}{rng.randint(10000000, 99999999)}"
    return [
        sn, extra_name or fake_name(rng), fake_date(rng, year, month, style),
        f"{rng.randint(100, 999)}-{rng.randint(1000, 9999)}-{rng.randint(100, 999)}",
        facility, rng.randint(1000, 9999),
        rng.randint(16, 44) if rng.random() < 0.6 else f"{rng.randint(16, 44)}YRS",
        f"{rng.randint(1, 60)} {rng.choice(['Palm', 'Cedar', 'Unity', 'Market', 'School'])} Road",
        int(phone) if phone_as_number else "0" + phone,
        rng.choice(CLIENT_TYPES), f"G{rng.randint(1, 5)}P{rng.randint(0, 4)}",
    ]


def write_workbook(path, rng, facility, year, month, n_rows, demo_name=None):
    wb = Workbook(write_only=True)
    if rng.random() < 0.03:  # cover sheet first, data on the second sheet
        cover = wb.create_sheet("Cover")
        cover.append(["Monthly summary form - see next sheet"])
    ws = wb.create_sheet("MSF")
    for _ in range(rng.choice([0, 0, 1, 2, 3, 5, 8])):  # title rows above the header
        ws.append([f"{facility} monthly summary form {MONTH_NAMES[month - 1]} {year}"])
        ws.append([None])
    keys = list(HEADER_VARIANTS)
    if rng.random() < 0.1:
        keys.remove("Parity")
    ws.append([rng.choice(HEADER_VARIANTS[k]) for k in keys])
    phone_as_number = rng.random() < 0.5
    demo_at = rng.randrange(n_rows) if demo_name else -1
    for i in range(n_rows):
        row = make_row(rng, i + 1, facility, year, month, phone_as_number,
                       demo_name if i == demo_at else None)
        ws.append(row[:len(keys)])
        if rng.random() < 0.02:
            ws.append([None] * len(keys))  # stray blank row
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="data/raw_xlsx")
    ap.add_argument("--files", type=int, default=544)
    ap.add_argument("--min-rows", type=int, default=80)
    ap.add_argument("--max-rows", type=int, default=450)
    ap.add_argument("--seed", type=int, default=42)
    a = ap.parse_args()
    rng = random.Random(a.seed)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)

    slots = [(f, y, m) for f in FACILITIES for y in range(2020, 2026) for m in range(1, 13)
             if (y, m) <= (2025, 3)]
    rng.shuffle(slots)
    slots = slots[:a.files]
    demo_slots = {slots[0]: DEMO_PEOPLE[0], slots[1]: DEMO_PEOPLE[0], slots[2]: DEMO_PEOPLE[0],
                  slots[3]: DEMO_PEOPLE[1]} if len(slots) >= 4 else {}
    for f, y, m in slots:
        suffix = rng.choices(["", " Updated", " - Copy"], [0.88, 0.07, 0.05])[0]
        name = f"{f} {MONTH_NAMES[m - 1]} {y} MSF{suffix}.xlsx"
        write_workbook(out / name, rng, f, y, m, rng.randint(a.min_rows, a.max_rows),
                       demo_slots.get((f, y, m)))
    print(f"Wrote {len(slots)} synthetic workbooks to {out}")


if __name__ == "__main__":
    main()
